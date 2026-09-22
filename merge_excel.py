"""Merge the first sheet of .xlsx files with identical headers."""
import argparse
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell


def merge(folder, output, dedupe=False):
    root, destination = Path(folder).resolve(), Path(output).resolve()
    if not root.is_dir():
        raise ValueError('Input folder does not exist')
    if destination.suffix.lower() != '.xlsx':
        raise ValueError('Output must end in .xlsx')
    if destination.exists():
        raise ValueError('Output already exists; choose a new filename')
    files = sorted(p for p in root.iterdir() if p.suffix.lower() == '.xlsx' and p.is_file() and not p.name.startswith('~$') and p.resolve() != destination)
    if not files:
        raise ValueError('No .xlsx input files found')
    header, rows, seen = None, [], set()
    for path in files:
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            sheet = book.worksheets[0]
            values = sheet.iter_rows(values_only=True)
            current = next(values, None)
            if not current or any(v is None for v in current) or len(set(current)) != len(current):
                raise ValueError(f'Empty or duplicate header: {path.name}')
            if header is None:
                header = current
            elif current != header:
                raise ValueError(f'Headers differ: {path.name}')
            for row in values:
                if all(v is None for v in row):
                    continue
                if any(isinstance(v, str) and v.startswith('=') for v in row):
                    raise ValueError(f'Formula found in {path.name}; convert formulas to values first')
                if not dedupe or row not in seen:
                    rows.append(row)
                    if dedupe:
                        seen.add(row)
        finally:
            book.close()
    if len(rows) + 1 > 1048576:
        raise ValueError('Too many rows for one Excel worksheet')
    result = Workbook(write_only=True)
    sheet = result.create_sheet('Merged')
    for row in [header, *rows]:
        cells = []
        for value in row:
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                cell.data_type = 's'
            cells.append(cell)
        sheet.append(cells)
    # Exclusive creation prevents accidental replacement of an existing workbook.
    with destination.open('xb') as stream:
        result.save(stream)
    result.close()
    return len(files), len(rows)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('folder')
    parser.add_argument('output')
    parser.add_argument('--dedupe', action='store_true')
    args = parser.parse_args()
    try:
        files, rows = merge(args.folder, args.output, args.dedupe)
        print(f'Merged {files} workbook(s), {rows} data row(s) -> {args.output}')
    except (OSError, ValueError) as exc:
        parser.exit(1, f'Error: {exc}\n')


if __name__ == '__main__':
    main()
