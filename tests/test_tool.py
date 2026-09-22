import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from merge_excel import merge

def workbook(path, rows):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    book.save(path)
    book.close()

class MergerTests(unittest.TestCase):
    def test_merge_dedupe(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            workbook(root / 'a.xlsx', [('Name', 'Count'), ('Pen', 2)])
            workbook(root / 'b.xlsx', [('Name', 'Count'), ('Pen', 2), ('Book', 1)])
            output = root / 'merged.xlsx'
            self.assertEqual(merge(root, output, True), (2, 2))
            book = load_workbook(output)
            self.assertEqual(list(book.active.values), [('Name', 'Count'), ('Pen', 2), ('Book', 1)])
            book.close()
            with self.assertRaises(ValueError):
                merge(root, output)
    def test_header_mismatch_no_output(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            workbook(root / 'a.xlsx', [('A',), (1,)])
            workbook(root / 'b.xlsx', [('B',), (2,)])
            with self.assertRaises(ValueError):
                merge(root, root / 'out.xlsx')
            self.assertFalse((root / 'out.xlsx').exists())
    def test_formula_rejected(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            workbook(root / 'a.xlsx', [('A',), ('=1+1',)])
            with self.assertRaises(ValueError):
                merge(root, root / 'out.xlsx')
    def test_empty_folder(self):
        with tempfile.TemporaryDirectory() as folder:
            with self.assertRaises(ValueError):
                merge(folder, Path(folder) / 'out.xlsx')
